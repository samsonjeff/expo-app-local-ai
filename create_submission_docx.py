import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Set 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(51, 65, 85) # Slate 700

    # Header / Title Banner
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    p_title.paragraph_format.space_before = Pt(0)
    run_badge = p_title.add_run("AppBuildersPH Hackathon 2026 • Local AI")
    run_badge.font.size = Pt(10)
    run_badge.font.bold = True
    run_badge.font.color.rgb = RGBColor(16, 185, 129) # Emerald Green

    p_main_title = doc.add_paragraph()
    p_main_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_main_title.paragraph_format.space_after = Pt(4)
    run_main = p_main_title.add_run("Official Project Submission Checklist")
    run_main.font.size = Pt(22)
    run_main.font.bold = True
    run_main.font.color.rgb = RGBColor(15, 23, 42) # Slate 900

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    run_sub = p_sub.add_run("Project: MaQui — Offline Local AI Quiz & Reviewer")
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(79, 70, 229) # Indigo

    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(16)
    p_div_run = p_div.add_run("―" * 58)
    p_div_run.font.color.rgb = RGBColor(226, 232, 240)

    def add_section_header(title, icon=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(f"{title}")
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)

    def add_field(label, value):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.space_before = Pt(2)
        r_label = p.add_run(f"•  {label}: ")
        r_label.font.bold = True
        r_label.font.color.rgb = RGBColor(30, 41, 59)
        r_val = p.add_run(value)
        r_val.font.color.rgb = RGBColor(51, 65, 85)

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        r_bold = p.add_run(bold_prefix)
        r_bold.font.bold = True
        r_bold.font.color.rgb = RGBColor(30, 41, 59)
        r_text = p.add_run(text)
        r_text.font.color.rgb = RGBColor(51, 65, 85)

    # ---------------- SECTION 1 ----------------
    add_section_header("1. The Project")
    add_field("Project Name", "MaQui (Offline Local AI Quiz & Reviewer)")
    add_field("Short Description", 
              "MaQui is a 100% offline, on-device AI-powered assessment and active-recall study companion for students and teachers. It ingests lecture notes and presentation slides (PDF, PPTX, DOCX, TXT) and autonomously generates curriculum-aligned mock examinations (Multiple Choice, True/False, Identification, and Essay) along with spaced-repetition flashcards (SRS)—running entirely on local device hardware with zero cloud latency, zero API costs, and zero data telemetry.")
    add_field("Team Members", "Rochelle, Sedrick, Jeff Samson")
    add_field("Public GitHub Repository", "https://github.com/samsonjeff/expo-app-local-ai (Branch: polishing/sedrick / main)")

    # ---------------- SECTION 2 ----------------
    add_section_header("2. The Proof")
    add_field("Demo Video", "[Insert your Google Drive / YouTube / Loom link here]")
    add_field("X / LinkedIn Video URL", "[Insert your X / LinkedIn post URL here]")

    p_run_loc = doc.add_paragraph()
    p_run_loc.paragraph_format.space_before = Pt(4)
    p_run_loc.paragraph_format.space_after = Pt(4)
    r_loc = p_run_loc.add_run("•  What Runs Locally:")
    r_loc.font.bold = True
    r_loc.font.color.rgb = RGBColor(30, 41, 59)

    add_bullet("Document Ingestion & Chunking: ", "Multi-format document text extraction (PDF via syncfusion_flutter_pdf, PowerPoint PPTX / Word DOCX via archive & XML stream parsing) running in dedicated Dart background Isolates.")
    add_bullet("LLM Inference: ", "On-device quantized GGUF neural network execution (e.g., Qwen 2.5 1.5B / Phi-4-mini / Llama 3.2 3B) running through llama.cpp native C++ bindings via dart:ffi.")
    add_bullet("JSON Schema Repair & Validation: ", "Real-time streaming token validation and deterministic syntax repair for structured question-answer outputs.")
    add_bullet("Assessment & Grading Engine: ", "Instant offline evaluation, student answer comparison, AI conceptual rationale generation, and grading rubrics.")
    add_bullet("Flashcard Spaced Repetition (SRS): ", "Local Leitner/SM-2 spaced repetition memory algorithm and review scheduling.")
    add_bullet("Local Database & Sandboxing: ", "Type-safe SQLite relational database (sqflite + drift) for storing quizzes, question banks, study attempts, and deck states.")
    add_bullet("Export & Print Formatting: ", "Offline generation and sharing of printable Student Exam Papers and Teacher Answer Keys with rubrics (pdf & printing).")

    p_run_net = doc.add_paragraph()
    p_run_net.paragraph_format.space_before = Pt(6)
    p_run_net.paragraph_format.space_after = Pt(4)
    r_net = p_run_net.add_run("•  What Requires Internet:")
    r_net.font.bold = True
    r_net.font.color.rgb = RGBColor(30, 41, 59)

    add_bullet("One-Time Model Download Only: ", "Fetching the quantized GGUF weights to device storage during initial setup from Hugging Face.")
    add_bullet("Zero Operational Internet: ", "Once downloaded, the application requires exactly 0 bytes of internet connectivity and operates completely in Airplane Mode (Wi-Fi and Cellular disabled).")

    # ---------------- SECTION 3 ----------------
    add_section_header("3. The Disclosures")
    
    p_m = doc.add_paragraph()
    p_m.paragraph_format.space_before = Pt(4)
    p_m.paragraph_format.space_after = Pt(4)
    r_m = p_m.add_run("•  Models Used:")
    r_m.font.bold = True
    r_m.font.color.rgb = RGBColor(30, 41, 59)

    add_bullet("Primary (4GB RAM tier devices): ", "Qwen2.5-1.5B-Instruct-Q4_K_M.gguf (capped to 4096 context tokens for guaranteed zero-OOM execution on budget smartphones).")
    add_bullet("Secondary (6GB–8GB RAM tier devices): ", "Phi-4-mini-Instruct-Q4_K_M.gguf / Llama-3.2-3B-Instruct-Q4_K_M.gguf (up to 8192 context window).")
    add_bullet("Dual-Mode Mock Fallback Engine: ", "Built-in fallback engine for instant multi-platform testing without native C++ compilation.")

    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(6)
    p_t.paragraph_format.space_after = Pt(4)
    r_t = p_t.add_run("•  Technologies and Frameworks:")
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(30, 41, 59)

    add_bullet("Mobile Framework: ", "Flutter (Dart 3.x), targeting mobile-first Android and iOS architectures.")
    add_bullet("State Management: ", "flutter_riverpod (Riverpod 2.x).")
    add_bullet("Inference Runtime: ", "llama.cpp compiled via CMake/Android NDK, bound via dart:ffi.")
    add_bullet("Database: ", "SQLite via sqflite and drift.")
    add_bullet("Document Parsing: ", "syncfusion_flutter_pdf, archive, xml.")
    add_bullet("Persistence & Services: ", "flutter_foreground_task, shared_preferences, path_provider.")
    add_bullet("Export & Sharing: ", "pdf, printing, share_plus.")
    add_bullet("UI & Animations: ", "Material Design 3, google_fonts (Inter), custom isolate-safe tactile micro-interaction system (TactilePressCard, FlipCard3D, AiPulseGlow, StaggeredEntranceItem).")

    add_field("APIs and Cloud Services", "None. Zero external AI APIs used (no OpenAI, no Anthropic, no Google Gemini API, no cloud servers, no proxy backends). Only standard public CDN (Hugging Face) for downloading public GGUF weights during setup.")
    add_field("Existing Code and Assets", "Open-source Flutter packages from pub.dev, open-source llama.cpp C++ engine, and custom MaQui application branding assets (assets/MaQui-light-mode.png, assets/MaQui-dark-mode.png).")
    add_field("AI Development Tools", "Google Antigravity / Gemini IDE for architectural pair-programming, test authoring, and UI polish.")

    # ---------------- SECTION 4 (KEY QUESTION) ----------------
    add_section_header("4. Key Requirement: Why does this product benefit from running AI locally?")

    # Callout Box / Table for Highlight
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)

    cell = table.cell(0, 0)
    set_cell_background(cell, "F0FDF4") # Light Green Shading
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    p_box = cell.paragraphs[0]
    p_box.paragraph_format.space_before = Pt(4)
    p_box.paragraph_format.space_after = Pt(4)
    r_box_title = p_box.add_run("Executive Summary:\n")
    r_box_title.font.bold = True
    r_box_title.font.size = Pt(11.5)
    r_box_title.font.color.rgb = RGBColor(22, 101, 52)
    r_box_desc = p_box.add_run(
        "Running AI locally is the essential technological foundation of MaQui. It guarantees 100% student and institutional privacy, eliminates recurring API costs for public education, ensures uncompromised offline availability in areas with weak or zero connectivity, and enables cheat-proof assessment integrity in Airplane Mode."
    )
    r_box_desc.font.size = Pt(10.5)
    r_box_desc.font.color.rgb = RGBColor(22, 101, 52)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Detailed Points
    def add_benefit_point(number, title, explanation):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        r_num = p.add_run(f"{number}. {title}\n")
        r_num.font.bold = True
        r_num.font.size = Pt(11.5)
        r_num.font.color.rgb = RGBColor(15, 23, 42)
        r_exp = p.add_run(explanation)
        r_exp.font.size = Pt(10.5)
        r_exp.font.color.rgb = RGBColor(51, 65, 85)

    add_benefit_point(
        "1",
        "100% Student & Institutional Data Privacy (Zero Cloud Telemetry)",
        "In educational settings, study materials frequently consist of unreleased teacher syllabi, unpublished research notes, internal lecture slides, and confidential midterm review items. Cloud-hosted LLMs transmit every word over third-party servers, where it can be logged, monetized, or ingested into training corpora. MaQui processes every single token in the local sandboxed memory of the student's personal phone, making compliance with privacy regulations (such as GDPR and the Philippine Data Privacy Act) absolute and verified by design."
    )

    add_benefit_point(
        "2",
        "Zero Recurring API Costs for Students and Public Schools",
        "Cloud AI applications impose ongoing operational costs: every 30-question quiz generated from a 15-page slide deck costs money per token on proprietary cloud APIs. For public school students and budget-conscious universities across developing regions, subscriptions or credit limits create a digital divide. Running inference on-device democratizes access: once installed, generating 1,000 quizzes costs ₱0.00 / $0.00 forever."
    )

    add_benefit_point(
        "3",
        "Bulletproof Accessibility in Low-Connectivity & Classroom Environments",
        "Internet connectivity across public transport, rural provinces, dormitories, and crowded campus Wi-Fi networks is notoriously unstable or expensive. Furthermore, educational institutions and testing centers frequently prohibit or restrict internet access in classrooms. Because MaQui operates entirely offline, students can study and generate fresh mock exams anywhere—whether riding a bus, studying during power blackouts, or reviewing in an airplane seat."
    )

    add_benefit_point(
        "4",
        "Deterministic Reliability Without Cloud Outages or Rate-Limits",
        "Cloud LLM APIs are subject to server degradation, queueing delays, HTTP 429 rate limits, and breaking API schema changes right before exam week. MaQui provides instantaneous local execution with zero network roundtrip latency and zero external points of failure."
    )

    add_benefit_point(
        "5",
        "Exam Integrity Under Airplane Mode",
        "When taking assessments or studying with active recall, educators can require students to toggle Airplane Mode on their devices. With MaQui, the AI tutor and quiz generator continue to work at peak capability while preventing students from browsing external cheat sheets or messaging classmates."
    )

    # Footer note
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(24)
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = p_foot.add_run("Submission Deadline: 10:00 AM, October 10 • AppBuildersPH Hackathon 2026")
    r_foot.font.size = Pt(9.5)
    r_foot.font.italic = True
    r_foot.font.color.rgb = RGBColor(148, 163, 184)

    output_path = r"c:\Users\Rochelle\expo-app-local-ai\AppBuildersPH_Hackathon_2026_Submission_MaQui.docx"
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    create_document()
